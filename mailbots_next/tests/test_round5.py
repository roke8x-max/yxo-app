"""Round-5 coverage: runtime MODE, live/test gates, split fan-out, real
inbound handler, digest boundary, notify fallback, DSK-T6.

Live-mode send assertions spy smtplib.SMTP_SSL (no network): the real
execute_action code path runs, only the socket is faked.
"""
import io
import json
from email.mime.base import MIMEBase
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email import encoders
from unittest.mock import Mock, MagicMock, patch

import pytest

from mailbots_next.config import (
    BOT_CONFIG_DB_PATH,
    DEDUP_DB_PATH,
    DRAFT_NUMS_DB_PATH,
    DAILY_COUNTERS_PATH,
    OPS_OWNER_EMAIL,
    EmailType,
)
from mailbots_next.core.dedup import init_db
from mailbots_next.core.store import (
    init_bot_config_db,
    get_bot_config_connection,
    seed_owner_mapping,
)
from mailbots_next.core.extract import ExtractedRow


@pytest.fixture(autouse=True)
def _r5_dbs():
    for p in (BOT_CONFIG_DB_PATH, DEDUP_DB_PATH, DRAFT_NUMS_DB_PATH, DAILY_COUNTERS_PATH):
        if p.exists():
            try:
                p.unlink()
            except PermissionError:
                pass
    init_db()
    init_bot_config_db()
    seed_owner_mapping()
    conn = get_bot_config_connection()
    try:
        for company, to in (("太平洋", ["tp-a@test.com"]), ("东盟", ["tp-b@test.com"])):
            conn.execute(
                """INSERT OR REPLACE INTO bot_config (bot, scope, key, to_addrs, cc_addrs, extra)
                   VALUES (?, 'company', ?, ?, '[]', '{}')""",
                ("all", company, json.dumps(to)),
            )
        conn.commit()
    finally:
        conn.close()
    try:
        yield
    finally:
        from mailbots_next.config.provider import invalidate_cache
        invalidate_cache()


def _harness():
    """Real MailProcessor on the temp seed DB (records stay real)."""
    from mailbots_next.serve import MailProcessor
    from mailbots_next.config import snapshot
    proc = MailProcessor(snapshot())
    proc.notifier = Mock()
    return proc


def _smtp_spy():
    """Fake SMTP_SSL class: real execute_action path, zero network."""
    server = Mock()
    server.login.return_value = None
    server.sendmail.return_value = {}
    ctx = MagicMock()
    ctx.__enter__.return_value = server
    cls = Mock(return_value=ctx)
    return cls, server


def _draft_raw(code, box):
    msg = MIMEMultipart()
    msg["Subject"] = f"草单 {code}"
    msg["From"] = "youlia@yxologistics.com"
    msg.attach(MIMEText("body", "plain"))
    att = MIMEBase("application", "pdf")
    att.set_payload(b"test")
    encoders.encode_base64(att)
    att.add_header("Content-Disposition", "attachment",
                   filename=f"{box}-260810-123456已加密.pdf")
    msg.attach(att)
    return msg.as_bytes()


# ---------------------------------------------------------- Sec 0 live flag
class TestRuntimeMode:
    def test_mode_is_runtime_dynamic(self, monkeypatch):
        import mailbots_next.config as cfg
        monkeypatch.setenv("MAILBOT_MODE", "live")
        assert cfg.is_live() is True
        monkeypatch.setenv("MAILBOT_MODE", "test")
        assert cfg.is_live() is False

    def test_live_flag_activates_send_path(self, monkeypatch):
        """Live 下可路由邮件驱动真实发送路径，sendmail 被调用."""
        monkeypatch.setenv("MAILBOT_MODE", "live")
        proc = _harness()
        raw = _draft_raw("CQWLJT260810001-SPB", "CICU1000001")
        cls, server = _smtp_spy()
        with patch("smtplib.SMTP_SSL", cls):
            proc.process_email("maoxiaoyang@cqtransit.com", "运单草单",
                               "r5-live-1", 0, "草单 CQWLJT260810001-SPB",
                               "youlia@yxologistics.com", "", raw)
        assert server.sendmail.call_count == 1


# ------------------------------------------------- Sec 1 live/test 门禁
class TestModeGates:
    def _dsk_row(self, box):
        return ExtractedRow(row_idx=0, email_type=EmailType.DSK, container_no=box)

    def _dsk_routing(self, company, box):
        from mailbots_next.core.routing import RoutingResult
        to = {"太平洋": ["tp-a@test.com"], "东盟": ["tp-b@test.com"]}[company]
        return RoutingResult(company, "m@x.com", to, [], "box_match", True)

    def test_live_mode_sends_and_writes(self, monkeypatch):
        """Live：多公司 DSK 真发送且真写库（wraps 观察 + 查库验证）."""
        import mailbots_next.core.act as act_mod
        from mailbots_next.core.decide import Decision
        monkeypatch.setenv("MAILBOT_MODE", "live")
        real_write = act_mod.write_dsk_timestamp
        calls = []
        import sqlite3

        def _spy_write(box, field, ts):
            calls.append((box, field))
            return real_write(box, field, ts)

        cls, server = _smtp_spy()
        with patch("smtplib.SMTP_SSL", cls), \
             patch.object(act_mod, "write_dsk_timestamp", side_effect=_spy_write):
            for box, company in (("CICU1000001", "太平洋"), ("CICU1000002", "东盟")):
                ok, detail = act_mod.execute_action(
                    Decision("T1", "forward", "t"), self._dsk_row(box),
                    self._dsk_routing(company, box), b"raw",
                    "maoxiaoyang@cqtransit.com", "pw", f"r5-livew-{box}")
                assert (ok, detail) == (True, "forwarded")
        assert server.sendmail.call_count == 2
        assert sorted(calls) == [("CICU1000001", "DSK"), ("CICU1000002", "DSK")]
        from mailbots_next.config import YXO_DB_PATH
        conn = sqlite3.connect(str(YXO_DB_PATH))
        try:
            got = {r[0]: r[1] for r in conn.execute(
                "SELECT 箱号, dsk FROM records WHERE 箱号 LIKE 'CICU100000%'").fetchall()}
        finally:
            conn.close()
        assert got.get("CICU1000001") and got.get("CICU1000002")

    def test_test_mode_does_not_send_or_write(self, monkeypatch):
        """Test：同样邮件不发送不写库，返回 test-mode-skipped."""
        import mailbots_next.core.act as act_mod
        from mailbots_next.core.decide import Decision
        monkeypatch.setenv("MAILBOT_MODE", "test")
        cls, server = _smtp_spy()
        with patch("smtplib.SMTP_SSL", cls), \
             patch.object(act_mod, "write_dsk_timestamp") as mock_write:
            ok, detail = act_mod.execute_action(
                Decision("T1", "forward", "t"),
                self._dsk_row("CICU1000001"),
                self._dsk_routing("太平洋", "CICU1000001"), b"raw",
                "maoxiaoyang@cqtransit.com", "pw", "r5-testw-1")
        assert (ok, detail) == (True, "test-mode-skipped")
        cls.assert_not_called()
        mock_write.assert_not_called()


# ------------------------------------------------- Sec 2 反泄漏拆分
def _waybill_msg():
    msg = MIMEMultipart()
    msg["Subject"] = "Waybill cover"
    msg["From"] = "docwbfb@yxologistics.com"
    msg.attach(MIMEText("cover text", "plain"))
    att = MIMEBase("application", "octet-stream")
    att.set_payload(b"info")
    encoders.encode_base64(att)
    att.add_header("Content-Disposition", "attachment", filename="info.txt")
    msg.attach(att)
    return msg


def _sent_parts(sendmail_mock):
    """Extract (recipients, text, files, subject) per sendmail call."""
    import email as email_mod
    out = []
    for call in sendmail_mock.call_args_list:
        _, to_addrs, msg_str = call.args
        m = email_mod.message_from_string(msg_str)
        texts, files = [], []
        for part in m.walk():
            if part.is_multipart():
                continue
            ctype = part.get_content_type()
            disp = str(part.get("Content-Disposition", "")).lower()
            if "attachment" in disp:
                files.append(part.get_filename())
            elif ctype in ("text/plain", "text/html"):
                payload = part.get_payload(decode=True)
                if payload:
                    texts.append(payload.decode("utf-8", errors="replace"))
        out.append({"to": list(to_addrs), "texts": texts,
                    "files": files, "subject": str(m["Subject"])})
    return out


def _wb_row(idx, code, box, rwb):
    return ExtractedRow(
        row_idx=idx, email_type=EmailType.WAYBILL,
        customer_code=code, container_no=box,
        raw_data={"waybill_row": {"客户编码": code, "箱号": box, "运单号": rwb}},
    )


class TestSplitFanout:
    def test_multicompany_waybill_sends_per_company_no_leak(self, monkeypatch):
        """2 公司 waybill → 2 次发送，各自文本只含本公司数据."""
        monkeypatch.setenv("MAILBOT_MODE", "live")
        proc = _harness()
        raw = _waybill_msg().as_bytes()
        rows = [
            _wb_row(0, "CQWLJT260810001-SPB", "CICU1000001", "RWB1"),
            _wb_row(1, "CQWLJT260810002-VXN", "CICU1000002", "RWB2"),
        ]
        cls, server = _smtp_spy()
        with patch("smtplib.SMTP_SSL", cls), \
             patch("mailbots_next.serve.extract_email", return_value=rows):
            proc.process_email("maoxiaoyang@cqtransit.com", "运单号",
                               "r5-split-1", 0, "Waybill cover",
                               "docwbfb@yxologistics.com", "", raw)
        assert server.sendmail.call_count == 2
        sent = _sent_parts(server.sendmail)
        by_to = {tuple(sorted(s["to"])): s for s in sent}
        assert ("tp-a@test.com",) in by_to and ("tp-b@test.com",) in by_to
        mail_a = by_to[("tp-a@test.com",)]
        mail_b = by_to[("tp-b@test.com",)]
        assert any("CICU1000001" in t for t in mail_a["texts"])
        assert any("CQWLJT260810001-SPB" in t for t in mail_a["texts"])
        assert not any("CICU1000002" in t or "CQWLJT260810002-VXN" in t
                       for t in mail_a["texts"])
        assert any("CICU1000002" in t for t in mail_b["texts"])
        assert not any("CICU1000001" in t or "CQWLJT260810001-SPB" in t
                       for t in mail_b["texts"])

    def test_single_company_still_one_send_full(self, monkeypatch):
        """单公司邮件 → 恰 1 次发送，内容与整封一致."""
        monkeypatch.setenv("MAILBOT_MODE", "live")
        proc = _harness()
        raw = _waybill_msg().as_bytes()
        rows = [_wb_row(0, "CQWLJT260810001-SPB", "CICU1000001", "RWB1")]
        cls, server = _smtp_spy()
        with patch("smtplib.SMTP_SSL", cls), \
             patch("mailbots_next.serve.extract_email", return_value=rows):
            proc.process_email("maoxiaoyang@cqtransit.com", "运单号",
                               "r5-single-1", 0, "Waybill cover",
                               "docwbfb@yxologistics.com", "", raw)
        assert server.sendmail.call_count == 1
        sent = _sent_parts(server.sendmail)[0]
        assert sent["subject"] == "Waybill cover"
        assert any("cover text" in t for t in sent["texts"])

    def test_rewrite_xls_filtered_keeps_only_requested(self):
        """附件过滤原语：重写后 xls 只含目标行（openpyxl 实字节往返）."""
        import openpyxl
        from mailbots_next.core.act import rewrite_xls_filtered
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["客户编码", "箱号", "运单号"])
        ws.append(["CQWLJT260810001-SPB", "CICU1000001", "RWB1"])
        ws.append(["CQWLJT260810002-VXN", "CICU1000002", "RWB2"])
        buf = io.BytesIO()
        wb.save(buf)
        new_bytes, new_name = rewrite_xls_filtered(
            buf.getvalue(), "运单号.xlsx",
            [{"客户编码": "CQWLJT260810001-SPB", "箱号": "CICU1000001", "运单号": "RWB1"}])
        assert new_bytes and new_name == "运单号.xlsx"
        wb2 = openpyxl.load_workbook(io.BytesIO(new_bytes), read_only=True)
        data_rows = list(wb2.active.iter_rows(values_only=True))[1:]
        assert len(data_rows) == 1
        flat = " ".join(str(c) for c in data_rows[0])
        assert "CICU1000001" in flat and "CICU1000002" not in flat

    def test_dsk_split_per_box_no_leak(self, monkeypatch):
        """同公司双箱 DSK → 逐箱 2 封：标题/正文/附件均按箱隔离."""
        monkeypatch.setenv("MAILBOT_MODE", "live")
        proc = _harness()
        html = ("<html><body><table>"
                "<tr><td>CICU1000001</td><td>A</td></tr>"
                "<tr><td>CICU1000002</td><td>B</td></tr>"
                "</table></body></html>")
        msg = MIMEMultipart()
        msg["Subject"] = "DSK notice"
        msg["From"] = "kasa@rtsb.de"
        msg.attach(MIMEText(html, "html"))
        for fn, body in (("CICU1000001.PDF", b"a"), ("CICU1000002.PDF", b"b"),
                         ("readme.txt", b"r")):
            att = MIMEBase("application", "octet-stream")
            att.set_payload(body)
            encoders.encode_base64(att)
            att.add_header("Content-Disposition", "attachment", filename=fn)
            msg.attach(att)
        raw = msg.as_bytes()
        cls, server = _smtp_spy()
        with patch("smtplib.SMTP_SSL", cls):
            proc.process_email("maoxiaoyang@cqtransit.com", "DSK",
                               "r5-dsk-1", 0, "DSK notice",
                               "kasa@rtsb.de", "", raw)
        assert server.sendmail.call_count == 2
        sent = _sent_parts(server.sendmail)
        by_box = {}
        for s in sent:
            key = "CICU1000001" if "CICU1000001" in s["subject"] else "CICU1000002"
            by_box[key] = s
        assert set(by_box) == {"CICU1000001", "CICU1000002"}
        for box, other in (("CICU1000001", "CICU1000002"),
                           ("CICU1000002", "CICU1000001")):
            s = by_box[box]
            assert box in s["subject"] and "CQWLJT" in s["subject"]
            assert not any(other in t for t in s["texts"])
            assert f"{box}.PDF" in s["files"]
            assert f"{other}.PDF" not in s["files"]
            assert "readme.txt" in s["files"]


# ------------------------------------------------- Sec 4 digest 月界
class TestDigestBoundary:
    def test_digest_next_time_month_boundary(self):
        from datetime import datetime
        from mailbots_next.serve import _next_digest_time
        assert _next_digest_time(datetime(2026, 1, 31, 10, 0)) == datetime(2026, 2, 1, 9, 0)
        assert _next_digest_time(datetime(2026, 3, 31, 23, 0)) == datetime(2026, 4, 1, 9, 0)
        assert _next_digest_time(datetime(2026, 2, 28, 10, 0)) == datetime(2026, 3, 1, 9, 0)
        assert _next_digest_time(datetime(2026, 5, 1, 8, 0)) == datetime(2026, 5, 1, 9, 0)


# ------------------------------------------------- Sec 5 企微降级
class TestNotifyFallback:
    def _notifier(self):
        from mailbots_next.core.notify import WeComNotifier
        n = WeComNotifier()
        n._notify_by_name = Mock(return_value=(True, "wxwork"))
        return n

    def test_notify_unknown_recipient_escalates_to_owner_email(self):
        n = self._notifier()
        cls, server = _smtp_spy()
        with patch("smtplib.SMTP_SSL", cls):
            assert n.notify("alarm", ["ghost@example.com"], "原始告警正文") is True
        assert server.sendmail.call_count == 1
        _, to_addrs, msg_str = server.sendmail.call_args.args
        assert to_addrs == [OPS_OWNER_EMAIL]
        import email as email_mod
        from email.header import decode_header
        parsed = email_mod.message_from_string(msg_str)
        body = parsed.get_payload(decode=True).decode("utf-8")
        assert "ghost@example.com" in body and "原始告警正文" in body
        subj, enc = decode_header(parsed["Subject"])[0]
        subject = subj.decode(enc or "utf-8") if isinstance(subj, bytes) else subj
        assert "企微未映射" in subject and "ghost@example.com" in subject
        n._notify_by_name.assert_not_called()

    def test_notify_mapped_recipient_still_wecom(self):
        n = self._notifier()
        cls, server = _smtp_spy()
        with patch("smtplib.SMTP_SSL", cls):
            assert n.notify("alarm", ["maoxiaoyang@cqtransit.com"], "t") is True
        n._notify_by_name.assert_called_once()
        cls.assert_not_called()


# ------------------------------------------------- Sec 6 DSK-T6
class TestDskReuseTier:
    def test_decide_dsk_reused_box_triggers_T6(self):
        from mailbots_next.core.extract import ExtractedRow
        from mailbots_next.core.routing import route_dsk
        from mailbots_next.core.decide import decide_dsk
        recs = [
            {"id": 1, "客户编码": "CQWLJT260810001-SPB", "箱号": "CICU1000001",
             "company": "太平洋", "状态": "正常", "is_deleted": 0},
            {"id": 2, "客户编码": "CQWLJT260810002-VXN", "箱号": "CICU1000001",
             "company": "东盟", "状态": "正常", "is_deleted": 0},
        ]
        row = ExtractedRow(row_idx=0, email_type=EmailType.DSK, container_no="CICU1000001")
        routing = route_dsk(row, recs)
        assert routing.route_source == "box_reuse"
        d = decide_dsk(row, routing, recs)
        assert d.tier == "T6" and d.action == "alarm"
